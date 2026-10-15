# -*- coding: utf-8 -*-
"""
Cocina (backend) de la aplicacion de certificados de Academia Horizonte.

Toda la logica de negocio vive aqui. La interfaz (Salon) solo muestra
resultados; no calcula ni decide.
"""

import os
from pathlib import Path

from flask import Flask, render_template, request
from openpyxl import load_workbook

APP = Flask(__name__)

# Ladrillo: rutas de la bodega de datos (inmutables)
BASE_DIR = Path(__file__).resolve().parent
INSUMOS_DIR = BASE_DIR / "insumos"
ARCHIVO_MAESTRO = INSUMOS_DIR / "Maestro_Estudiantes.xlsx"
# Compatibilidad: si existe un nombre sin "(1)", usarlo; si no, usar el archivo real.
ARCHIVO_EVALUACIONES = (
    INSUMOS_DIR / "Registro_Evaluaciones.xlsx"
    if (INSUMOS_DIR / "Registro_Evaluaciones.xlsx").exists()
    else INSUMOS_DIR / "Registro_Evaluaciones (1).xlsx"
)


# ============================================================
# LADRILLOS REUTILIZADOS DE LA LOGICA YA VALIDADA
# ============================================================

def leer_hoja_xlsx(ruta, nombre_hoja=None):
    """Lee una hoja de Excel y devuelve (encabezados, filas)."""
    # Ladrillo: abrir el libro con openpyxl (solo valores)
    libro = load_workbook(ruta, data_only=True)
    # Ladrillo: elegir hoja por nombre si existe; si no, la primera
    if nombre_hoja and nombre_hoja in libro.sheetnames:
        hoja = libro[nombre_hoja]
    else:
        hoja = libro[libro.sheetnames[0]]

    # Ladrillo: iterar filas y convertir cada una a lista
    filas = list(hoja.iter_rows(values_only=True))
    if not filas:
        return [], []
    encabezados = [str(valor).strip() if valor is not None else "" for valor in filas[0]]
    datos = [list(fila) for fila in filas[1:]]
    return encabezados, datos


def calcular_promedio(valores):
    """Calcula el promedio de una lista; devuelve 0.0 si esta vacia."""
    # Ladrillo: longitud
    cantidad = len(valores)
    # Ladrillo: condicion para evitar division por cero
    if cantidad == 0:
        return 0.0
    # Ladrillo: suma y division
    total = sum(valores)
    return total / cantidad


def decidir_certificado(promedio, asistencia):
    """Aplica las reglas de negocio de AGENTS.md."""
    # Ladrillo: condicion 1 (limite incluido)
    if asistencia < 80:
        return "Sin certificado"
    # Ladrillo: condicion 2 (limite incluido)
    if promedio >= 70:
        return "Aprobacion"
    # Ladrillo: caso restante
    return "Participacion"


def preparar_datos():
    """
    Cruza maestro y evaluaciones, calcula resultados e inconsistencias.

    Devuelve un diccionario listo para la plantilla.
    """
    # Ladrillo: leer maestro y evaluaciones
    enc_m, filas_m = leer_hoja_xlsx(ARCHIVO_MAESTRO, "Estudiantes 1")
    enc_e, filas_e = leer_hoja_xlsx(ARCHIVO_EVALUACIONES, "Evaluaciones 1")

    # Ladrillo: mapear encabezados a indices
    idx_m = {nombre: i for i, nombre in enumerate(enc_m)}
    idx_e = {nombre: i for i, nombre in enumerate(enc_e)}

    # Ladrillo: indexar maestro por (Identificacion, Programa)
    maestro_por_clave = {}
    for fila in filas_m:
        clave = (str(fila[idx_m["Identificacion"]]).strip(),
                 str(fila[idx_m["Programa"]]).strip())
        maestro_por_clave[clave] = {
            "identificacion": clave[0],
            "nombre": str(fila[idx_m["Nombre_Completo"]]).strip(),
            "correo": str(fila[idx_m["Correo"]]).strip(),
            "programa": clave[1],
            "cohorte": str(fila[idx_m["Cohorte"]]).strip(),
        }

    # Ladrillo: agrupar evaluaciones por (Identificacion, Programa)
    evaluaciones_por_clave = {}
    for fila in filas_e:
        clave = (str(fila[idx_e["Identificacion"]]).strip(),
                 str(fila[idx_e["Programa"]]).strip())
        evaluaciones_por_clave.setdefault(clave, []).append({
            "modulo": str(fila[idx_e["Modulo"]]).strip(),
            "nota": float(fila[idx_e["Nota"]]),
            "asistencia": float(fila[idx_e["Asistencia_Pct"]]),
            "fecha_cierre": str(fila[idx_e["Fecha_Cierre"]]).strip(),
        })

    resultados = []
    # Ladrillo: bucle principal sobre cada estudiante/programa evaluado
    for clave, lista_eval in evaluaciones_por_clave.items():
        # Solo certificamos estudiantes registrados en el maestro
        if clave not in maestro_por_clave:
            continue

        # Ladrillo: extraer listas para calcular promedios
        notas = [e["nota"] for e in lista_eval]
        asistencias = [e["asistencia"] for e in lista_eval]
        modulos_cursados = len(lista_eval)

        # Ladrillo: aplicar funciones
        promedio_nota = calcular_promedio(notas)
        promedio_asistencia = calcular_promedio(asistencias)
        certificado = decidir_certificado(promedio_nota, promedio_asistencia)

        maestro = maestro_por_clave[clave]
        resultados.append({
            "programa": clave[1],
            "nombre": maestro["nombre"],
            "identificacion": clave[0],
            "correo": maestro["correo"],
            "cohorte": maestro["cohorte"],
            "modulos_cursados": modulos_cursados,
            "modulos": ", ".join(e["modulo"] for e in lista_eval),
            "promedio_nota": round(promedio_nota, 2),
            "promedio_asistencia": round(promedio_asistencia, 2),
            "certificado": certificado,
            "clase_certificado": {
                "Aprobacion": "cert-aprobacion",
                "Participacion": "cert-participacion",
                "Sin certificado": "cert-sin",
            }[certificado],
        })

    # Ladrillo: ordenar por programa y nombre
    resultados.sort(key=lambda r: (r["programa"], r["nombre"], r["identificacion"]))

    # --- Inconsistencias ---
    claves_maestro = set(maestro_por_clave.keys())
    claves_eval = set(evaluaciones_por_clave.keys())

    # 1) Maestro sin evaluaciones
    sin_evaluaciones = []
    for clave in sorted(claves_maestro - claves_eval):
        m = maestro_por_clave[clave]
        sin_evaluaciones.append({
            "identificacion": m["identificacion"],
            "nombre": m["nombre"],
            "programa": m["programa"],
        })

    # 2) Evaluaciones no presentes en el maestro
    fuera_de_maestro = []
    for clave in sorted(claves_eval - claves_maestro):
        lista = evaluaciones_por_clave[clave]
        notas = [e["nota"] for e in lista]
        asistencias = [e["asistencia"] for e in lista]
        fuera_de_maestro.append({
            "identificacion": clave[0],
            "programa": clave[1],
            "modulos": len(lista),
            "promedio_nota": round(calcular_promedio(notas), 2),
            "promedio_asistencia": round(calcular_promedio(asistencias), 2),
            "certificado": decidir_certificado(calcular_promedio(notas),
                                               calcular_promedio(asistencias)),
        })

    # 3) Modulos incompletos (menos de 4)
    modulos_incompletos = []
    for clave, lista in sorted(evaluaciones_por_clave.items()):
        if len(lista) < 4:
            ident = clave[0]
            nombre_maestro = maestro_por_clave.get(clave, {}).get("nombre", "NO REGISTRADO")
            modulos_incompletos.append({
                "identificacion": ident,
                "nombre": nombre_maestro,
                "programa": clave[1],
                "cantidad": len(lista),
                "modulos": ", ".join(e["modulo"] for e in lista),
            })

    # Ladrillo: contadores generales (solo grupos certificables)
    contadores = {
        "aprobacion": sum(1 for r in resultados if r["certificado"] == "Aprobacion"),
        "participacion": sum(1 for r in resultados if r["certificado"] == "Participacion"),
        "sin_certificado": sum(1 for r in resultados if r["certificado"] == "Sin certificado"),
        "total": len(resultados),
    }

    programas = sorted(set(m["programa"] for m in maestro_por_clave.values()))

    return {
        "resultados": resultados,
        "contadores": contadores,
        "programas": programas,
        "sin_evaluaciones": sin_evaluaciones,
        "fuera_de_maestro": fuera_de_maestro,
        "modulos_incompletos": modulos_incompletos,
        "archivo_evaluaciones": ARCHIVO_EVALUACIONES.name,
    }


@APP.route("/")
def index():
    datos = preparar_datos()
    # Ladrillo: filtro opcional por programa
    programa_filtro = request.args.get("programa", "").strip()
    resultados = datos["resultados"]
    if programa_filtro:
        resultados = [r for r in resultados if r["programa"] == programa_filtro]

    return render_template(
        "index.html",
        resultados=resultados,
        contadores=datos["contadores"],
        programas=datos["programas"],
        programa_filtro=programa_filtro,
        inconsistencias={
            "sin_evaluaciones": datos["sin_evaluaciones"],
            "fuera_de_maestro": datos["fuera_de_maestro"],
            "modulos_incompletos": datos["modulos_incompletos"],
        },
    )


if __name__ == "__main__":
    # Ladrillo: arranque del servidor en modo desarrollo
    print("Academia Horizonte - Servidor iniciado en http://localhost:5000")
    APP.run(host="0.0.0.0", port=5000, debug=True)
PY